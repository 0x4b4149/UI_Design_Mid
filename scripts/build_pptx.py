"""Build a visual PPTX backup from the report PDF pages.

The PDF is the primary submission. This PPTX mirrors it page-for-page so the
deck can be presented even without a browser. To change text, edit
report_meta.json / report_template.html and rebuild both files.
"""

import sys
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Inches

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
PDF = ROOT / "docs" / "report" / "Skyline_Weather_期中專題報告.pdf"
PPTX = ROOT / "docs" / "report" / "Skyline_Weather_期中專題報告.pptx"
TMP = ROOT / "docs" / "report" / "preview" / "pptx_pages"
TMP.mkdir(parents=True, exist_ok=True)


def main():
    doc = pymupdf.open(PDF)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for index, page in enumerate(doc):
        pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
        image_path = TMP / f"page-{index + 1:02d}.png"
        pix.save(image_path)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height
        )

    props = prs.core_properties
    props.title = "Skyline Weather｜期中專題報告"
    props.subject = "從 weather.gov Bad UI 調查到 Django + Neon 全端實作"
    props.author = "Skyline Weather Team"
    props.comments = "Produced from the HTML report deck; PDF is the primary submission."
    prs.save(PPTX)
    print(f"PPTX: {PPTX}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
